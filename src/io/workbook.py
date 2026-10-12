
"""Workbook loading and validation for Project BlindSpot."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


REQUIRED_COLUMNS = ["Vendor ID", "Vendor Name"]


def find_header_row(file_path, sheet_name):
    """Find the row containing the vendor inventory column headers."""

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' was not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        sheet = workbook[sheet_name]

        for row_number, row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1,
        ):
            values = {
                str(value).strip().casefold()
                for value in row
                if value is not None
            }

            if "vendor id" in values and "vendor name" in values:
                return row_number

        raise ValueError(
            f"Could not find the vendor header row in '{sheet_name}'."
        )

    finally:
        workbook.close()


def load_inventory(file_path, sheet_name="Vendor Inventory"):
    """
    Load and validate the vendor inventory.

    Returns:
        inventory: DataFrame containing the source vendor records.
        messages: List of validation messages.
        metadata: Workbook and validation details.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Workbook not found: {path.resolve()}"
        )

    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(
            "Please provide an Excel workbook (.xlsx or .xlsm)."
        )

    header_row = find_header_row(path, sheet_name)

    inventory = pd.read_excel(
        path,
        sheet_name=sheet_name,
        header=header_row - 1,
        dtype=object,
        engine="openpyxl",
    )

    # Remove only completely empty columns created by spreadsheet layout.
    inventory = inventory.dropna(axis=1, how="all")

    # Normalize header whitespace without changing source values.
    canonical_headers = {
        "vendor id": "Vendor ID",
        "vendor name": "Vendor Name",
    }
    inventory.columns = [
        canonical_headers.get(
            str(column).strip().casefold(),
            str(column).strip(),
        )
        for column in inventory.columns
    ]

    messages = []

    missing_columns = [
        column
        for column in REQUIRED_COLUMNS
        if column not in inventory.columns
    ]

    if missing_columns:
        raise ValueError(
            "Missing required columns: "
            + ", ".join(missing_columns)
        )

    # Preserve all source rows, including any supplied worked example.
    inventory = inventory.reset_index(drop=True)

    blank_ids = inventory["Vendor ID"].isna()
    blank_names = inventory["Vendor Name"].isna()

    if blank_ids.any():
        messages.append(
            f"{int(blank_ids.sum())} row(s) have no Vendor ID."
        )

    if blank_names.any():
        messages.append(
            f"{int(blank_names.sum())} row(s) have no Vendor Name."
        )

    duplicate_ids = (
        inventory["Vendor ID"]
        .dropna()
        .astype(str)
        .str.strip()
        .str.casefold()
        .duplicated(keep=False)
    )

    if duplicate_ids.any():
        messages.append(
            "Duplicate Vendor ID values were found. "
            "Review them before scoring."
        )

    if inventory.empty:
        messages.append("The vendor inventory contains no data rows.")

    if not messages:
        messages.append(
            "Workbook validation passed: required columns are present."
        )

    metadata = {
        "source_file": path.name,
        "sheet_name": sheet_name,
        "header_row": header_row,
        "vendor_rows": len(inventory),
        "column_count": len(inventory.columns),
        "columns": list(inventory.columns),
        "validation_passed": not any(
            "missing" in message.lower()
            or "duplicate" in message.lower()
            for message in messages
        ),
    }

    return inventory, messages, metadata
