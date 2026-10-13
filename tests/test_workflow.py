"""Focused regression tests for the vendor evidence workflow."""

from io import BytesIO
import ast
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from src.assessment.classifier import (
    CONFIRMED,
    INCONCLUSIVE,
    NOT_DETECTED,
    classify_ai_usage,
)
from src.assessment.evidence import grade_evidence
from src.assessment.evidence import create_evidence_record
from src.assessment.criticality import initialize_inventory_criticality
from src.assessment.risk import (
    calculate_ai_cyber_risk_score,
    create_risk_finding,
)
from src.assessment.scores import calculate_evidence_score
from src.collection.discovery import normalize_url
from src.collection.extractor import extract_evidence_candidates
from src.collection.fetcher import _is_safe_public_url
from src.collection.fetcher import PublicPageFetcher
from src.collection.fetcher import FetchResult
from src.collection.investigation import (
    discover_github_repositories,
    discover_relevant_link_records,
    discover_relevant_links,
    run_full_investigation,
)
from src.io.validation import duplicate_vendor_ids, validate_inventory
from src.io.workbook import load_inventory
from src.ui.report_export import build_assessment_workbook
from src.ui.collection_page import (
    _determine_source_type,
    merge_evidence_records,
    show_collection_page,
)
from src.ui.intelligence_mesh import (
    build_relationship_graph,
    effective_relationship_status,
)
from src.ui.review_centre import (
    CORROBORATION_OPTIONS,
    RECENCY_OPTIONS,
    RELEVANCE_OPTIONS,
    REVIEW_REASONS,
)


class TestEvidenceWorkflow(unittest.TestCase):
    def test_aggregate_evidence_score_uses_only_accepted_relevant_grades(self):
        score = calculate_evidence_score([
            {
                "review_status": "Accepted",
                "grade": {"score": 5, "service_relevance": "relevant"},
            },
            {
                "review_status": "Accepted",
                "grade": {"score": 3, "service_relevance": "unclear"},
            },
            {
                "review_status": "Needs review",
                "grade": {"score": 5, "service_relevance": "relevant"},
            },
            {
                "review_status": "Accepted",
                "grade": {"score": 5, "service_relevance": "not_relevant"},
            },
        ])
        self.assertEqual(score["score"], 80.0)
        self.assertEqual(score["evidence_count"], 2)

    def test_ai_cyber_risk_score_requires_accepted_same_vendor_evidence(self):
        finding = {
            "finding_id": "RF-1",
            "vendor_id": "V-1",
            "severity": "Critical",
            "likelihood": "Likely",
            "status": "Needs review",
            "evidence_ids": ["E-1"],
        }
        unsupported = calculate_ai_cyber_risk_score(
            [finding],
            evidence_records=[{
                "vendor_id": "V-2",
                "evidence_id": "E-1",
                "review_status": "Accepted",
            }],
        )
        self.assertIsNone(unsupported["score"])
        self.assertEqual(unsupported["unsupported_findings"], 1)

        scored = calculate_ai_cyber_risk_score(
            [finding],
            evidence_records=[{
                "vendor_id": "V-1",
                "evidence_id": "E-1",
                "review_status": "Accepted",
            }],
        )
        self.assertEqual(scored["score"], 100.0)

    def test_live_investigation_collects_traceable_candidate_records(self):
        class FakeFetcher:
            def fetch_page(self, url):
                return FetchResult(
                    url=url,
                    status="success",
                    retrieved_at="2026-10-12T00:00:00+00:00",
                    http_status=200,
                    content_type="text/html",
                    title="Vendor technical post",
                    html=(
                        "<html><body><a href='/careers'>Careers</a>"
                        "<p>Our generative AI assistant uses a model.</p>"
                        "</body></html>"
                    ),
                )

        progress = []
        result = run_full_investigation(
            vendor_id="V-1",
            vendor_name="Example Vendor",
            seed_urls=["https://vendor.example"],
            max_pages=1,
            include_github=False,
            fetcher=FakeFetcher(),
            progress_callback=lambda **event: progress.append(event),
        )
        self.assertEqual(result["pages_attempted"], 1)
        self.assertEqual(result["pages_fetched"], 1)
        self.assertEqual(len(result["evidence_records"]), 1)
        evidence = result["evidence_records"][0]
        self.assertEqual(evidence["vendor_id"], "V-1")
        self.assertEqual(evidence["review_status"], "Needs review")
        self.assertEqual(evidence["data_origin"], "live_public_web")
        self.assertTrue(evidence["source_url"].startswith("https://"))
        self.assertTrue(evidence["excerpt"])
        self.assertEqual(
            result["collection_log"][0]["data_origin"],
            "live_public_web",
        )
        self.assertTrue(progress)
        self.assertFalse(result["search_scope_complete"])

    def test_link_discovery_stays_with_vendor_domain_or_public_github(self):
        urls = discover_relevant_links(
            """
            <a href="https://docs.vendor.example/ai">Docs</a>
            <a href="https://attacker.example/careers">Careers</a>
            <a href="https://github.com/vendor/repo">Public repository</a>
            """,
            "https://vendor.example/",
            {"vendor.example"},
        )
        self.assertIn("https://docs.vendor.example/ai", urls)
        self.assertIn("https://github.com/vendor/repo", urls)
        self.assertNotIn("https://attacker.example/careers", urls)

    def test_link_discovery_records_crawl_provenance(self):
        records = discover_relevant_link_records(
            """
            <a href="/docs/ai">AI docs</a>
            <a href="https://github.com/vendor/repo">Public repository</a>
            """,
            "https://vendor.example/blog/post",
            {"vendor.example"},
        )
        self.assertEqual(records[0]["url"], "https://vendor.example/docs/ai")
        self.assertEqual(
            records[0]["from_url"],
            "https://vendor.example/blog/post",
        )
        self.assertFalse(records[0]["is_public_repository"])
        self.assertTrue(records[1]["is_public_repository"])

    def test_live_investigation_honors_crawl_depth(self):
        class FakeFetcher:
            def fetch_page(self, url):
                html = "<html><body><a href='/docs/ai'>AI docs</a></body></html>"
                if url.endswith("/docs/ai"):
                    html = (
                        "<html><body><p>Our generative AI assistant helps "
                        "analysts.</p></body></html>"
                    )
                return FetchResult(
                    url=url,
                    status="success",
                    retrieved_at="2026-10-12T00:00:00+00:00",
                    http_status=200,
                    content_type="text/html",
                    title="Vendor page",
                    html=html,
                )

        result = run_full_investigation(
            vendor_id="V-1",
            vendor_name="Example Vendor",
            seed_urls=["https://vendor.example"],
            max_pages=2,
            max_depth=1,
            include_github=False,
            fetcher=FakeFetcher(),
        )
        linked = [
            item for item in result["candidates"]
            if item["discovery_method"] == "linked_public_page_candidate"
        ]
        self.assertEqual(result["crawl_depth_limit"], 1)
        self.assertGreaterEqual(result["link_candidates_discovered"], 1)
        self.assertEqual(linked[0]["crawl_depth"], 1)
        self.assertEqual(linked[0]["parent_url"], "https://vendor.example/")

    def test_github_search_uses_bounded_public_api(self):
        response = MagicMock()
        response.status_code = 200
        response.headers = {
            "Content-Type": "application/json",
            "Content-Length": "100",
        }
        response.iter_content.return_value = [
            json.dumps({
                "items": [{
                    "html_url": "https://github.com/example/project",
                }]
            }).encode()
        ]
        with patch(
            "src.collection.investigation.requests.get",
            return_value=response,
        ) as request:
            urls, note = discover_github_repositories("Example Vendor")

        self.assertEqual(urls, ["https://github.com/example/project"])
        self.assertIn("candidate", note)
        self.assertFalse(request.call_args.kwargs["allow_redirects"])
        self.assertEqual(request.call_args.kwargs["timeout"], (4, 10))

    def test_grade_is_not_a_probability_and_requires_valid_url(self):
        grade = grade_evidence(
            "official_product_documentation",
            "https://vendor.example/product",
            "direct",
            "relevant",
            "current",
            "independent",
        )

        self.assertEqual(grade.score, 5)
        self.assertIn("proposed points", grade.rationale)
        self.assertIsNone(
            grade_evidence(
                "official_product_documentation",
                "file:///private",
                "direct",
                "relevant",
            ).score
        )

    def test_classifier_requires_accepted_direct_relevant_evidence(self):
        accepted = {
            "evidence_id": "E-1",
            "review_status": "Accepted",
            "grade": {
                "score": 5,
                "directness": "direct",
                "service_relevance": "relevant",
                "recency": "current",
            },
        }
        self.assertEqual(
            classify_ai_usage([accepted]).status,
            CONFIRMED,
        )

        rejected = dict(accepted, review_status="Rejected")
        self.assertEqual(
            classify_ai_usage([rejected], collection_complete=True).status,
            NOT_DETECTED,
        )

    def test_dated_direct_evidence_cannot_establish_confirmed_status(self):
        evidence = {
            "evidence_id": "E-DATED",
            "review_status": "Accepted",
            "grade": {
                "score": 5,
                "directness": "direct",
                "service_relevance": "relevant",
                "recency": "dated",
            },
        }
        self.assertEqual(classify_ai_usage([evidence]).status, "Suspected")

    def test_classifier_does_not_skip_pending_or_incomplete_scope(self):
        accepted = {
            "evidence_id": "E-1",
            "review_status": "Accepted",
            "grade": {
                "score": 5,
                "directness": "direct",
                "service_relevance": "relevant",
                "recency": "current",
            },
        }
        pending = {
            "evidence_id": "E-2",
            "review_status": "Needs review",
            "grade": {},
        }

        self.assertEqual(
            classify_ai_usage([accepted, pending]).status,
            INCONCLUSIVE,
        )
        self.assertEqual(classify_ai_usage([]).status, INCONCLUSIVE)
        self.assertEqual(
            classify_ai_usage([], collection_complete=True).status,
            NOT_DETECTED,
        )

    def test_extraction_records_candidate_only_with_stable_source_id(self):
        fetch_result = {
            "status": "success",
            "url": "https://vendor.example/docs",
            "title": "Product docs",
            "retrieved_at": "2026-10-12T00:00:00+00:00",
            "html": "<html><body><p>Our generative AI assistant helps teams.</p></body></html>",
        }
        first = extract_evidence_candidates(fetch_result)
        second = extract_evidence_candidates(fetch_result)

        self.assertEqual(len(first), 1)
        self.assertEqual(first[0].evidence_id, second[0].evidence_id)
        self.assertEqual(first[0].review_status, "Needs review")
        self.assertEqual(first[0].source_url, fetch_result["url"])

    def test_deduplication_preserves_existing_review_and_grade(self):
        existing = {
            "evidence_id": "E-1",
            "vendor_id": "V-1",
            "review_status": "Accepted",
            "reviewer": "Analyst",
            "review_notes": "Verified against source context.",
            "data_origin": "live_public_web",
            "discovery_method": "linked_public_page_candidate",
            "run_id": "RUN-1",
            "grade": {"score": 5, "directness": "direct"},
        }
        recaptured = {
            "evidence_id": "E-1",
            "vendor_id": "V-1",
            "review_status": "Needs review",
            "reviewer": "",
            "grade": {"score": 2},
        }

        merged, added, duplicates = merge_evidence_records(
            [existing], [recaptured]
        )

        self.assertEqual((added, duplicates), (0, 1))
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0]["review_status"], "Accepted")
        self.assertEqual(merged[0]["review_notes"], "Verified against source context.")
        self.assertEqual(merged[0]["grade"]["score"], 5)
        self.assertEqual(merged[0]["data_origin"], "live_public_web")
        self.assertEqual(
            merged[0]["discovery_method"],
            "linked_public_page_candidate",
        )
        self.assertEqual(merged[0]["run_id"], "RUN-1")

    def test_url_normalization_rejects_invalid_ports(self):
        self.assertIsNone(normalize_url("https://vendor.example:bad/path"))

    def test_review_options_match_the_grade_api(self):
        for relevance in RELEVANCE_OPTIONS:
            for recency in RECENCY_OPTIONS:
                for corroboration in CORROBORATION_OPTIONS:
                    grade = grade_evidence(
                        "official_product_documentation",
                        "https://vendor.example/docs",
                        "unclear",
                        relevance,
                        recency,
                        corroboration,
                    )
                    self.assertIsNotNone(grade.score)

        self.assertTrue(REVIEW_REASONS)

    def test_evidence_ids_are_stable_and_vendor_scoped(self):
        candidate = {
            "evidence_id": "E-ABC",
            "source_url": "https://vendor.example/docs",
            "excerpt": "AI assistant product detail",
        }
        first = create_evidence_record(
            candidate,
            "V-1",
            "Vendor One",
            "official_product_documentation",
        )
        repeated = create_evidence_record(
            candidate,
            "V-1",
            "Vendor One",
            "official_product_documentation",
        )
        other_vendor = create_evidence_record(
            candidate,
            "V-2",
            "Vendor Two",
            "official_product_documentation",
        )

        self.assertEqual(first.evidence_id, repeated.evidence_id)
        self.assertNotEqual(first.evidence_id, other_vendor.evidence_id)
        self.assertEqual(first.review_status, "Needs review")

    def test_fetcher_blocks_non_public_destinations(self):
        self.assertFalse(_is_safe_public_url("http://127.0.0.1/"))
        self.assertFalse(_is_safe_public_url("http://169.254.169.254/"))
        self.assertFalse(_is_safe_public_url("http://10.0.0.1/"))
        self.assertFalse(_is_safe_public_url("https://8.8.8.8:8443/"))
        self.assertTrue(_is_safe_public_url("https://8.8.8.8/"))

    def test_dns_names_resolving_to_private_addresses_are_blocked(self):
        with patch(
            "src.collection.fetcher.socket.getaddrinfo",
            return_value=[
                (2, 1, 6, "", ("192.168.1.5", 443)),
            ],
        ):
            self.assertFalse(_is_safe_public_url("https://vendor.example/"))

    def test_redirect_to_private_destination_is_blocked(self):
        fetcher = PublicPageFetcher(delay_seconds=0)
        fetcher._robots_cache["https://example.com/robots.txt"] = True
        response = MagicMock()
        response.is_redirect = True
        response.headers = {"Location": "http://127.0.0.1/admin"}
        fetcher.session.get = MagicMock(return_value=response)

        with patch(
            "src.collection.fetcher._is_safe_public_url",
            side_effect=[True, True, True, False],
        ):
            result = fetcher.fetch_page("https://example.com/")

        fetcher.close()
        self.assertEqual(result.status, "blocked")
        self.assertEqual(result.error_category, "unsafe_redirect")
        fetcher.session.get.assert_called_once()

    def test_risk_findings_keep_vendor_rationale_and_review_metadata(self):
        finding = create_risk_finding(
            vendor_id="V-1",
            vendor_name="Example Vendor",
            theme="evidence_gaps",
            rationale="No public assurance evidence was collected.",
            severity="Unassessed",
            evidence_ids=["E-1"],
            reviewer="Reviewer",
            reviewed_at="2026-10-12T00:00:00+00:00",
        )

        self.assertEqual(finding.vendor_id, "V-1")
        self.assertEqual(finding.evidence_ids, ["E-1"])
        self.assertEqual(finding.reviewer, "Reviewer")
        self.assertEqual(finding.status, "Needs review")


class TestInventoryAndExport(unittest.TestCase):
    def test_inventory_upload_criticality_uses_explicit_rating_columns(self):
        inventory = pd.DataFrame([{
            "Vendor ID": "V-1",
            "Vendor Name": "Example",
            "Data Sensitivity Rating": 5,
            "Operational Dependency Rating": 4,
            "Regulatory Exposure Rating": 3,
            "Payment Involvement Rating": 2,
            "Annual Data Volume Rating": 1,
        }])
        assessment = initialize_inventory_criticality(inventory)["V-1"]
        self.assertEqual(assessment["result"].score, 70.0)
        self.assertEqual(assessment["result"].coverage_percent, 100.0)
        self.assertTrue(assessment["auto_calculated"])

    def test_inventory_criticality_does_not_invent_numeric_volume_bands(self):
        inventory = pd.DataFrame([{
            "Vendor ID": "V-1",
            "Vendor Name": "Example",
            "Data Sensitivity": "High",
            "Operational Dependency": "Critical",
            "Regulatory Exposure": "High",
            "Payment Involvement": "Yes",
            "Annual Data Volume": 50000,
        }])
        result = initialize_inventory_criticality(inventory)["V-1"]["result"]
        self.assertIsNone(result.score)
        self.assertIn("annual_data_volume", result.missing_factors)

    def test_vendor_score_export_keeps_three_scores_separate(self):
        inventory = pd.DataFrame([{
            "Vendor ID": "V-1",
            "Vendor Name": "Example",
            "Data Sensitivity Rating": 5,
            "Operational Dependency Rating": 4,
            "Regulatory Exposure Rating": 3,
            "Payment Involvement Rating": 2,
            "Annual Data Volume Rating": 1,
        }])
        criticality = initialize_inventory_criticality(inventory)
        evidence = [{
            "vendor_id": "V-1",
            "evidence_id": "E-1",
            "review_status": "Accepted",
            "grade": {
                "score": 4,
                "service_relevance": "relevant",
            },
        }]
        finding = {
            "vendor_id": "V-1",
            "finding_id": "RF-1",
            "severity": "High",
            "likelihood": "Likely",
            "status": "Needs review",
            "evidence_ids": ["E-1"],
        }
        state = {
            "inventory_df": inventory,
            "evidence_records": evidence,
            "review_log": [],
            "risk_findings": [finding],
            "collection_log": [],
            "criticality_ratings": criticality,
            "investigation_assessments": {},
            "relationship_records": [],
        }
        with patch("streamlit.session_state", state):
            exported = build_assessment_workbook()

        workbook = load_workbook(BytesIO(exported.getvalue()), read_only=True)
        scores = workbook["Vendor Scores"]
        columns = {cell.value: cell.column for cell in scores[1]}
        row = {
            name: scores.cell(2, column).value
            for name, column in columns.items()
        }
        self.assertEqual(row["Business Criticality (0-100)"], 70.0)
        self.assertEqual(row["Evidence Quality (0-100)"], 80.0)
        self.assertEqual(row["AI Cybersecurity Risk (0-100)"], 75.0)
        self.assertNotEqual(
            row["Business Criticality (0-100)"],
            row["AI Cybersecurity Risk (0-100)"],
        )
        workbook.close()

    def test_every_sidebar_route_has_one_page_branch(self):
        app_tree = ast.parse((PROJECT_ROOT / "app.py").read_text(encoding="utf-8"))
        route = next(
            node for node in app_tree.body
            if isinstance(node, ast.If)
            and isinstance(node.test, ast.Compare)
            and isinstance(node.test.left, ast.Name)
            and node.test.left.id == "selected_page"
        )

        labels = []
        current = route
        while isinstance(current, ast.If):
            if (
                isinstance(current.test, ast.Compare)
                and current.test.comparators
                and isinstance(current.test.comparators[0], ast.Constant)
            ):
                labels.append(current.test.comparators[0].value)
            current = current.orelse[0] if current.orelse else None

        expected = {
            "Overview",
            "Vendor Inventory",
            "AI Investigation",
            "Live Vendor Investigation Lab",
            "Public Evidence Collection",
            "Evidence Register",
            "Review Centre",
            "Investigation Summary",
            "Risk Assessment",
            "Intelligence Mesh",
            "Vendor Profile",
            "Reports & Export",
        }
        self.assertEqual(set(labels), expected)
        self.assertEqual(len(labels), len(set(labels)))

    def test_inventory_validation_catches_duplicates_and_missing_values(self):
        inventory = pd.DataFrame(
            {
                "Vendor ID": [" v-1 ", "V-1", None],
                "Vendor Name": ["One", "Two", " "],
                "Website URL": ["vendor.example", "bad url", ""],
            }
        )

        report = validate_inventory(inventory)
        codes = {issue.code for issue in report.issues}

        self.assertIn("DUPLICATE_VENDOR_ID", codes)
        self.assertIn("MISSING_VENDOR_ID", codes)
        self.assertIn("MISSING_VENDOR_NAME", codes)
        self.assertIn("INVALID_URL_FORMAT", codes)
        self.assertEqual(len(inventory), report.total_rows)
        self.assertEqual(duplicate_vendor_ids(inventory), ["v-1"])

    def test_workbook_loader_finds_header_and_rejects_missing_sheet(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Vendor Inventory"
        sheet.append(["Inventory extract"])
        sheet.append(["vendor id", "vendor name", "Website"])
        sheet.append(["V-1", "Example", "https://example.com"])

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "inventory.xlsx"
            workbook.save(path)
            loaded, _, metadata = load_inventory(path)
            self.assertEqual(loaded.iloc[0]["Vendor ID"], "V-1")
            self.assertEqual(metadata["header_row"], 2)
            with self.assertRaises(ValueError):
                load_inventory(path, sheet_name="Missing")
            corrupt_path = Path(directory) / "corrupt.xlsx"
            corrupt_path.write_text("not an Excel workbook")
            with self.assertRaises(BadZipFile):
                load_inventory(corrupt_path)

    def test_export_has_all_sheets_and_does_not_mutate_inventory(self):
        inventory = pd.DataFrame(
            {"Vendor ID": ["V-1"], "Vendor Name": ["Example"]}
        )
        original = inventory.copy(deep=True)
        state = {
            "inventory_df": inventory,
            "evidence_records": [
                {
                    "evidence_id": "E-1",
                    "vendor_id": "V-1",
                    "grade": {
                        "score": 4,
                        "rationale": "Reviewer explanation",
                    },
                }
            ],
            "review_log": [],
            "risk_findings": [],
            "collection_log": [],
            "criticality_ratings": {},
            "investigation_assessments": {},
        }

        with patch("streamlit.session_state", state):
            exported = build_assessment_workbook()

        workbook = load_workbook(BytesIO(exported.getvalue()), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Vendor Inventory",
                "Evidence Register",
                "Review Log",
                "Risk Findings",
                "Collection Log",
                "Criticality Assessments",
                "AI Investigation Summary",
                "Assessment Overview",
                "Intelligence Mesh",
                "Vendor Scores",
            ],
        )
        self.assertEqual(
            workbook["Evidence Register"]["C1"].value,
            "Evidence Grade - score",
        )
        self.assertEqual(
            workbook["Evidence Register"]["D1"].value,
            "Evidence Grade - rationale",
        )
        workbook.close()
        pd.testing.assert_frame_equal(inventory, original)

    def test_empty_export_still_contains_all_sheets(self):
        state = {
            "inventory_df": None,
            "evidence_records": [],
            "review_log": [],
            "risk_findings": [],
            "collection_log": [],
            "criticality_ratings": {},
            "investigation_assessments": {},
        }

        with patch("streamlit.session_state", state):
            exported = build_assessment_workbook()

        workbook = load_workbook(BytesIO(exported.getvalue()), read_only=True)
        self.assertEqual(len(workbook.sheetnames), 10)
        workbook.close()

    def test_mesh_status_requires_accepted_evidence_for_the_same_vendor(self):
        relationship = {
            "vendor_id": "V-1",
            "reviewer_status": "Confirmed",
            "status": "Needs review",
            "evidence_id": "E-1",
        }

        self.assertEqual(
            effective_relationship_status(
                relationship,
                [{
                    "vendor_id": "V-1",
                    "evidence_id": "E-1",
                    "review_status": "Needs review",
                }],
            ),
            "Needs review",
        )
        self.assertEqual(
            effective_relationship_status(
                relationship,
                [{
                    "vendor_id": "V-2",
                    "evidence_id": "E-1",
                    "review_status": "Accepted",
                }],
            ),
            "Needs review",
        )
        self.assertEqual(
            effective_relationship_status(
                relationship,
                [{
                    "vendor_id": "V-1",
                    "evidence_id": "E-1",
                    "review_status": "Accepted",
                }],
            ),
            "Confirmed",
        )
        self.assertEqual(
            effective_relationship_status(
                {
                    "vendor_id": "V-1",
                    "reviewer_status": "Confirmed",
                    "status": "Confirmed",
                },
                [],
            ),
            "Needs review",
        )
        self.assertEqual(
            effective_relationship_status(
                {
                    "vendor_id": "V-1",
                    "reviewer_status": "Hypothesis — needs review",
                    "status": "Hypothesis — needs review",
                },
                [],
            ),
            "Hypothesis — needs review",
        )

    def test_mesh_graph_uses_vendor_ids_and_separates_entity_types(self):
        inventory = pd.DataFrame({
            "Vendor ID": ["V-1", "V-2"],
            "Vendor Name": ["Shared display name", "Shared display name"],
        })
        relationships = [
            {
                "relationship_id": "REL-1",
                "vendor_id": "V-1",
                "vendor_name": "Shared display name",
                "target_name": "Example",
                "target_type": "AI Provider",
                "relationship_type": "Uses AI service",
                "reviewer_status": "Confirmed",
                "evidence_id": "E-1",
            },
            {
                "relationship_id": "REL-2",
                "vendor_id": "V-2",
                "vendor_name": "Shared display name",
                "target_name": "Example",
                "target_type": "Cloud Provider",
                "relationship_type": "Uses cloud service",
                "reviewer_status": "Confirmed",
                "evidence_id": "E-2",
            },
        ]
        graph = build_relationship_graph(
            inventory,
            relationships,
            [
                {
                    "vendor_id": "V-1",
                    "evidence_id": "E-1",
                    "review_status": "Accepted",
                },
                {
                    "vendor_id": "V-2",
                    "evidence_id": "E-2",
                    "review_status": "Rejected",
                },
            ],
        )

        self.assertIn("vendor:v-1", graph)
        self.assertIn("vendor:v-2", graph)
        self.assertEqual(graph.number_of_nodes(), 4)
        self.assertEqual(graph.number_of_edges(), 2)
        statuses = {
            attributes["evidence_id"]: attributes["status"]
            for _, _, attributes in graph.edges(data=True)
        }
        self.assertEqual(statuses, {"E-1": "Confirmed", "E-2": "Needs review"})

    def test_mesh_relationships_are_exported_with_effective_status(self):
        state = {
            "inventory_df": pd.DataFrame(
                {"Vendor ID": ["V-1"], "Vendor Name": ["Example"]}
            ),
            "evidence_records": [],
            "review_log": [],
            "risk_findings": [],
            "collection_log": [],
            "criticality_ratings": {},
            "investigation_assessments": {},
            "relationship_records": [{
                "relationship_id": "REL-1",
                "vendor_id": "V-1",
                "vendor_name": "Example",
                "target_name": "AI Service",
                "target_type": "AI Provider",
                "relationship_type": "Uses AI service",
                "reviewer_status": "Confirmed",
                "status": "Confirmed",
                "evidence_id": "E-MISSING",
                "rationale": "Source is awaiting review.",
            }],
        }

        with patch("streamlit.session_state", state):
            exported = build_assessment_workbook()

        workbook = load_workbook(BytesIO(exported.getvalue()), read_only=True)
        self.assertEqual(workbook.sheetnames[8], "Intelligence Mesh")
        mesh = workbook["Intelligence Mesh"]
        self.assertEqual(mesh["A1"].value, "relationship_id")
        headings = {
            cell.value: cell.column
            for cell in mesh[1]
        }
        self.assertEqual(
            mesh.cell(2, headings["status"]).value,
            "Needs review",
        )
        workbook.close()

    def test_collection_page_source_type_determination(self):
        self.assertEqual(
            _determine_source_type(
                {"category": "product_documentation"},
                "https://vendor.example/docs",
            ),
            "official_product_documentation",
        )
        self.assertEqual(
            _determine_source_type(
                {"category": "engineering_and_jobs"},
                "https://vendor.example/careers",
            ),
            "relevant_technical_job_posting",
        )
        self.assertEqual(
            _determine_source_type(
                {},
                "https://vendor.example/blog/ai-feature",
            ),
            "official_technical_announcement",
        )
        self.assertEqual(
            _determine_source_type(
                {},
                "https://github.com/vendor/repo",
            ),
            "developer_reference",
        )
        self.assertEqual(
            _determine_source_type(
                {},
                "https://vendor.example/about",
            ),
            "indirect_technical_signal",
        )

    def test_classify_preliminary_ai_safeguards(self):
        from src.assessment.classifier import (
            CONFIRMED,
            SUSPECTED,
            NOT_DETECTED,
            INCONCLUSIVE,
            classify_preliminary_ai_usage,
        )

        # 1. Direct official product documentation supports Confirmed
        direct_doc = {
            "evidence_id": "E-DOC",
            "source_url": "https://vendor.example/docs/ai",
            "grade": {
                "score": 5,
                "source_type": "official_product_documentation",
                "directness": "direct",
                "service_relevance": "relevant",
                "recency": "current",
            },
        }
        res_confirmed = classify_preliminary_ai_usage([direct_doc], collection_complete=True)
        self.assertEqual(res_confirmed.status, CONFIRMED)
        self.assertEqual(res_confirmed.confidence, "Moderate")

        # 2. Safeguard: Technical job posting alone does NOT support Confirmed (capped at Suspected)
        job_posting = {
            "evidence_id": "E-JOB",
            "source_url": "https://vendor.example/careers/ai-engineer",
            "grade": {
                "score": 3,
                "source_type": "relevant_technical_job_posting",
                "directness": "direct",
                "service_relevance": "relevant",
                "recency": "current",
            },
        }
        res_job = classify_preliminary_ai_usage([job_posting], collection_complete=True)
        self.assertEqual(res_job.status, SUSPECTED)
        self.assertTrue(any("job posting alone" in r.lower() for r in res_job.rationale))

        # 3. Safeguard: Generic marketing alone does not establish AI use
        marketing = {
            "evidence_id": "E-MKT",
            "source_url": "https://vendor.example/about",
            "grade": {
                "score": 1,
                "source_type": "generic_marketing",
                "directness": "indirect",
                "service_relevance": "unclear",
                "recency": "current",
            },
        }
        res_mkt = classify_preliminary_ai_usage([marketing], collection_complete=True)
        self.assertEqual(res_mkt.status, NOT_DETECTED)

        # 4. Safeguard: Collection failure produces Inconclusive
        res_failed = classify_preliminary_ai_usage([], collection_complete=False, collection_failed=True, failure_reason="robots block")
        self.assertEqual(res_failed.status, INCONCLUSIVE)
        self.assertTrue(any("robots" in r.lower() or "restricted" in r.lower() for r in res_failed.rationale))

        # 5. Completed scope without qualifying evidence produces Not Detected
        res_not_detected = classify_preliminary_ai_usage([], collection_complete=True, collection_failed=False)
        self.assertEqual(res_not_detected.status, NOT_DETECTED)

    def test_pipeline_assess_single_vendor(self):
        from src.collection.pipeline import assess_single_vendor

        row = {
            "Vendor ID": "V-002",
            "Vendor Name": "Fiserv, Inc.",
            "Category": "Core Platform",
            "Public Website": "www.fiserv.com",
            "Service / Product Provided": "DNA core banking platform",
            "Operational Dependency": "Critical",
            "Data Classification Accessed": "Full customer master, balances and transactions for 1.4 million customers.",
            "Data Volume (annual)": "240 million transactions",
        }
        res = assess_single_vendor(row, run_live_osint=False)

        self.assertEqual(res["vendor_id"], "V-002")
        self.assertEqual(res["criticality_result"].tier, "Critical")
        self.assertEqual(res["depth_config"]["label"], "Full review — tier-driven")
        self.assertEqual(res["ai_assessment"].status, "Confirmed")
        self.assertGreater(len(res["evidence_records"]), 0)
        self.assertGreater(len(res["risk_findings"]), 0)
        self.assertIn("Criticality Tier", res["assessment_fields"])
        self.assertEqual(res["assessment_fields"]["Criticality Tier"], "Critical")
        self.assertEqual(res["assessment_fields"]["AI Usage Detected (Y/N)"], "Yes")
        self.assertIn("Azure", res["assessment_fields"]["Named AI Sub-processors / Model Providers"])

    def test_batch_automated_pipeline_execution_and_isolation(self):
        from src.io.workbook import load_inventory
        from src.collection.pipeline import run_batch_automated_pipeline

        df, _, _ = load_inventory("req/Meridian_Vendor_Input.xlsx")
        batch_out = run_batch_automated_pipeline(df, run_live_osint=False)

        enriched = batch_out["enriched_inventory_df"]
        self.assertEqual(len(enriched), 7)

        # Verify V-000 preserved verbatim
        v000 = enriched.iloc[0]
        self.assertEqual(v000["Vendor ID"], "V-000")
        self.assertEqual(v000["Criticality Tier"], "Critical")
        self.assertIn("Larkspur", str(v000["Vendor Name"]))

        # Verify all non-example rows have populated assessment fields
        for idx in range(1, len(enriched)):
            r = enriched.iloc[idx]
            self.assertIsNotNone(r["Criticality Tier"])
            self.assertIn(r["Criticality Tier"], {"Critical", "High", "Medium", "Low"})
            ai_flag = r.get("AI Usage Detected (Y/N)") or r.get("AI Usage Detected")
            self.assertIn(ai_flag, {"Yes", "No", "Inconclusive"})
            self.assertIsNotNone(r.get("AI Security Risk Classification"))

        # Verify failure isolation
        corrupt_df = df.copy()
        corrupt_df.loc[1, "Vendor Name"] = None  # test resilience
        isolated_out = run_batch_automated_pipeline(corrupt_df, run_live_osint=False)
        self.assertEqual(len(isolated_out["enriched_inventory_df"]), 7)

    def test_verify_exported_workbook_validation(self):
        from src.io.workbook import load_inventory
        from src.collection.pipeline import run_batch_automated_pipeline
        from src.ui.report_export import build_assessment_workbook, verify_exported_workbook

        df, _, _ = load_inventory("req/Meridian_Vendor_Input.xlsx")
        batch_out = run_batch_automated_pipeline(df, run_live_osint=False)

        session_state = {
            "inventory_df": batch_out["enriched_inventory_df"],
            "evidence_records": batch_out["all_evidence_records"],
            "review_log": [],
            "risk_findings": batch_out["all_risk_findings"],
            "collection_log": batch_out["all_collection_log"],
            "criticality_ratings": batch_out["criticality_ratings"],
            "investigation_assessments": batch_out["investigation_assessments"],
            "relationship_records": batch_out["all_relationship_records"],
        }

        with patch("streamlit.session_state", session_state):
            buf = build_assessment_workbook()

        ver_res = verify_exported_workbook(buf)
        self.assertTrue(ver_res["is_valid"])
        self.assertEqual(ver_res["sheet_count"], 10)
        self.assertEqual(ver_res["vendor_rows"], 7)
        self.assertTrue(ver_res["v000_preserved"])
        self.assertEqual(ver_res["populated_l_to_v_count"], 7)
        self.assertEqual(ver_res["missing_sheets"], [])


class TestAssessmentConsistencyAndIntegrity(unittest.TestCase):
    """Regression test suite for assessment consistency, risk claim gating, and evidence integrity."""

    @classmethod
    def setUpClass(cls):
        from src.io.workbook import load_inventory
        from src.collection.pipeline import run_batch_automated_pipeline

        cls.df, _, _ = load_inventory("req/Meridian_Vendor_Input.xlsx")
        cls.batch_out = run_batch_automated_pipeline(cls.df, run_live_osint=False)

    def _get_exported_workbook(self, session_overrides=None):
        from src.ui.report_export import build_assessment_workbook

        session_state = {
            "inventory_df": self.batch_out["enriched_inventory_df"].copy(),
            "evidence_records": list(self.batch_out["all_evidence_records"]),
            "review_log": [],
            "risk_findings": list(self.batch_out["all_risk_findings"]),
            "collection_log": list(self.batch_out["all_collection_log"]),
            "criticality_ratings": dict(self.batch_out["criticality_ratings"]),
            "investigation_assessments": dict(self.batch_out["investigation_assessments"]),
            "relationship_records": list(self.batch_out["all_relationship_records"]),
        }
        if session_overrides:
            session_state.update(session_overrides)

        with patch("streamlit.session_state", session_state):
            buf = build_assessment_workbook()
        return load_workbook(BytesIO(buf.getvalue()), data_only=True)

    def test_1_cross_sheet_ai_classification_consistency(self):
        """P0: Verify AI classifications match across Vendor Inventory, AI Investigation Summary, and Vendor Scores."""
        wb = self._get_exported_workbook()

        # Read Vendor Inventory rows
        inv_sheet = wb["Vendor Inventory"]
        inv_headers = [cell.value for cell in inv_sheet[1]]
        vid_col = inv_headers.index("Vendor ID") + 1
        ai_col = inv_headers.index("AI Usage Detected (Y/N)") + 1
        inv_ai = {}
        for r in range(2, inv_sheet.max_row + 1):
            vid = inv_sheet.cell(r, vid_col).value
            if vid:
                inv_ai[str(vid).strip()] = inv_sheet.cell(r, ai_col).value

        # Read AI Investigation Summary
        ai_sheet = wb["AI Investigation Summary"]
        ai_headers = [cell.value for cell in ai_sheet[1]]
        ai_vid_col = ai_headers.index("vendor_id") + 1
        ai_status_col = ai_headers.index("status") + 1
        summary_status = {}
        for r in range(2, ai_sheet.max_row + 1):
            vid = ai_sheet.cell(r, ai_vid_col).value
            if vid:
                summary_status[str(vid).strip()] = ai_sheet.cell(r, ai_status_col).value

        # Read Vendor Scores
        scores_sheet = wb["Vendor Scores"]
        score_headers = [cell.value for cell in scores_sheet[1]]
        sc_vid_col = score_headers.index("Vendor ID") + 1
        sc_ai_col = score_headers.index("AI-use Classification") + 1
        scores_ai = {}
        for r in range(2, scores_sheet.max_row + 1):
            vid = scores_sheet.cell(r, sc_vid_col).value
            if vid:
                scores_ai[str(vid).strip()] = scores_sheet.cell(r, sc_ai_col).value

        # V-001 AutomWorx: Inventory = No, Investigation Summary = Not Detected, Vendor Scores = Not Detected
        self.assertEqual(inv_ai.get("V-001"), "No")
        self.assertEqual(summary_status.get("V-001"), "Not Detected")
        self.assertEqual(scores_ai.get("V-001"), "Not Detected")
        self.assertNotIn("Suspected", [inv_ai.get("V-001"), summary_status.get("V-001"), scores_ai.get("V-001")])

        # V-003 FSSI: Inventory = Inconclusive, Investigation Summary = Inconclusive, Vendor Scores = Inconclusive
        self.assertEqual(inv_ai.get("V-003"), "Inconclusive")
        self.assertIn("Inconclusive", str(summary_status.get("V-003")))
        self.assertIn("Inconclusive", str(scores_ai.get("V-003")))
        self.assertNotIn("Suspected", [inv_ai.get("V-003"), summary_status.get("V-003"), scores_ai.get("V-003")])

        # V-004 Terrapin: Inventory = Inconclusive, Investigation Summary = Inconclusive, Vendor Scores = Inconclusive
        self.assertEqual(inv_ai.get("V-004"), "Inconclusive")
        self.assertIn("Inconclusive", str(summary_status.get("V-004")))
        self.assertIn("Inconclusive", str(scores_ai.get("V-004")))
        self.assertNotIn("Suspected", [inv_ai.get("V-004"), summary_status.get("V-004"), scores_ai.get("V-004")])

        # V-002, V-005, V-006: Inventory = Yes, Investigation Summary = Confirmed, Vendor Scores = Confirmed
        for vid in ["V-002", "V-005", "V-006"]:
            self.assertEqual(inv_ai.get(vid), "Yes")
            self.assertEqual(summary_status.get(vid), "Confirmed")
            self.assertEqual(scores_ai.get(vid), "Confirmed")
        wb.close()

    def test_2_ai_specific_risk_gating_for_no_and_inconclusive(self):
        """P0: Gating ensures No/Inconclusive AI classifications do not assert affirmative AI data processing claims."""
        findings = self.batch_out["all_risk_findings"]

        # V-001 AutomWorx (No / Not Detected):
        v1_findings = [f for f in findings if f.get("vendor_id") == "V-001"]
        self.assertTrue(len(v1_findings) >= 1)
        for f in v1_findings:
            desc = str(f.get("description", "")).lower()
            rat = str(f.get("rationale", "")).lower()
            self.assertNotIn("embedded ml models process live transaction", desc)
            self.assertNotIn("prompt exposure risk", desc)
            self.assertIn(f.get("risk_type"), {"assurance_gap", "conditional"})
            self.assertEqual(f.get("severity"), "Low")

        # V-003 FSSI & V-004 Terrapin (Inconclusive):
        for vid in ["V-003", "V-004"]:
            v_findings = [f for f in findings if f.get("vendor_id") == vid]
            self.assertTrue(len(v_findings) >= 1)
            for f in v_findings:
                self.assertIn(f.get("risk_type"), {"assurance_gap", "conditional"})
                desc = str(f.get("description", "")).lower()
                rat = str(f.get("rationale", "")).lower()
                self.assertNotIn("public evidence confirms ai use", rat)

    def test_3_curated_evidence_never_misrepresented_as_live_verified(self):
        """P0: Curated disclosures preserve provenance and are never represented as live-verified."""
        evidence = self.batch_out["all_evidence_records"]
        col_log = self.batch_out["all_collection_log"]

        curated_ev = [e for e in evidence if e.get("data_origin") == "curated_case_study_disclosure"]
        self.assertTrue(len(curated_ev) >= 1)
        for e in curated_ev:
            self.assertEqual(e.get("data_origin"), "curated_case_study_disclosure")
            self.assertIn("[Curated Reference]", str(e.get("source_title", "")))
            self.assertEqual(e.get("grade", {}).get("corroboration"), "not_corroborated")

        # Collection log records curated origin
        curated_log = [l for l in col_log if l.get("data_origin") == "curated_case_study_disclosure"]
        self.assertTrue(len(curated_log) >= 1)
        for l in curated_log:
            self.assertEqual(l.get("status"), "curated_disclosure")

    def test_4_provider_extraction_requiring_supporting_evidence(self):
        """P1: Provider extraction must be grounded in source text. No direct OpenAI inferred from Azure OpenAI."""
        from src.collection.extractor import extract_disclosed_model_providers

        # Fiserv DNA excerpt does not name Azure AI
        fiserv_excerpt = "Fiserv DNA integrates predictive analytics and machine learning capabilities for transaction pattern analysis."
        fiserv_providers = extract_disclosed_model_providers([{"excerpt": fiserv_excerpt}])
        self.assertNotIn("Microsoft Azure AI", fiserv_providers)
        self.assertNotIn("Azure", [p.lower() for p in fiserv_providers])

        # BNY excerpt mentions Azure OpenAI Service: direct OpenAI must NOT be extracted
        bny_excerpt = "BNY Enterprise AI Hub leverages Microsoft Azure OpenAI Service and custom machine learning models."
        bny_providers = extract_disclosed_model_providers([{"excerpt": bny_excerpt}])
        self.assertIn("Microsoft Azure OpenAI Service", bny_providers)
        self.assertNotIn("OpenAI", bny_providers)

        # Verify Fiserv Azure relationship in batch output is unverified
        relationships = self.batch_out["all_relationship_records"]
        v2_rels = [r for r in relationships if r.get("vendor_id") == "V-002" and "azure" in r.get("entity_name", "").lower()]
        for r in v2_rels:
            self.assertIn("unverified", r.get("status", "").lower())

    def test_5_duplicate_and_irrelevant_source_handling(self):
        """P1: Repeated passages from same URL deduplicated; third-party GitHub searches marked not relevant."""
        from src.collection.extractor import extract_evidence_candidates
        from src.collection.fetcher import FetchResult

        repeated_html = """
        <html><body>
        <p>Our platform uses machine learning algorithms for automated decisioning.</p>
        <p>Our platform uses machine learning algorithms for automated decisioning.</p>
        <p>Third paragraph with machine learning algorithms.</p>
        </body></html>
        """
        page = FetchResult(
            url="https://github.com/example/repo",
            status="success",
            retrieved_at="2026-10-13T00:00:00Z",
            http_status=200,
            html=repeated_html,
        )
        candidates = extract_evidence_candidates(page)
        self.assertEqual(len(candidates), 1)

    def test_6_reviewer_decisions_remaining_distinct_from_automated_findings(self):
        """P1: Reviewer actions record distinct decisions and synchronize dependent outputs."""
        from src.assessment.review import sync_vendor_after_review

        inventory = self.batch_out["enriched_inventory_df"].copy()
        evidence = [dict(e) for e in self.batch_out["all_evidence_records"]]
        investigations = dict(self.batch_out["investigation_assessments"])
        risks = [dict(f) for f in self.batch_out["all_risk_findings"]]

        # Simulate reviewer approving evidence for V-005
        for e in evidence:
            if e.get("vendor_id") == "V-005":
                e["review_status"] = "Accepted"

        res = sync_vendor_after_review(
            "V-005",
            evidence_records=evidence,
            inventory_df=inventory,
            investigation_assessments=investigations,
            risk_findings=risks,
            reviewer_name="Senior Risk Assessor",
        )
        self.assertEqual(res["vendor_id"], "V-005")
        self.assertEqual(res["detected_flag"], "Yes")

        v5_row = inventory[inventory["Vendor ID"] == "V-005"].iloc[0]
        self.assertIn("Reviewer (Senior Risk Assessor)", str(v5_row.get("Assessedd By / Date")))

    def test_7_excel_and_csv_exports_reflect_same_authoritative_assessment(self):
        """P1: Excel and CSV exports produce identical vendor assessment fields."""
        from src.ui.report_export import populate_export_inventory

        enriched_inv = self.batch_out["enriched_inventory_df"].copy()
        csv_df = populate_export_inventory(
            enriched_inv,
            assessments=self.batch_out["criticality_ratings"],
            investigations=self.batch_out["investigation_assessments"],
            evidence=self.batch_out["all_evidence_records"],
            findings=self.batch_out["all_risk_findings"],
            relationships=self.batch_out["all_relationship_records"],
        )

        wb = self._get_exported_workbook()
        inv_sheet = wb["Vendor Inventory"]
        inv_headers = [cell.value for cell in inv_sheet[1]]

        for r_idx in range(2, inv_sheet.max_row + 1):
            vid = inv_sheet.cell(r_idx, inv_headers.index("Vendor ID") + 1).value
            if not vid or str(vid).strip().lower() == "nan":
                continue
            vid_str = str(vid).strip()
            csv_row = csv_df[csv_df["Vendor ID"] == vid_str].iloc[0]

            for col in ["Criticality Tier", "AI Usage Detected (Y/N)", "AI Security Risk Classification"]:
                if col in inv_headers and col in csv_df.columns:
                    excel_val = inv_sheet.cell(r_idx, inv_headers.index(col) + 1).value
                    csv_val = csv_row[col]
                    self.assertEqual(
                        str(excel_val or "").strip(),
                        str(csv_val or "").strip(),
                        f"Mismatch for vendor {vid_str} in column {col}: Excel={excel_val}, CSV={csv_val}"
                    )
        wb.close()

    def test_8_preservation_of_input_workbook_and_v000_example(self):
        """P1: req/Meridian_Vendor_Input.xlsx and V-000 Larkspur worked example preserved verbatim."""
        from src.io.workbook import load_inventory

        raw_df, _, _ = load_inventory("req/Meridian_Vendor_Input.xlsx")
        v0_raw = raw_df[raw_df["Vendor ID"] == "V-000"].iloc[0]
        self.assertIn("Larkspur", str(v0_raw["Vendor Name"]))

        enriched = self.batch_out["enriched_inventory_df"]
        v0_enriched = enriched[enriched["Vendor ID"] == "V-000"].iloc[0]
        self.assertEqual(v0_enriched["Vendor ID"], "V-000")
        self.assertEqual(v0_enriched["Criticality Tier"], "Critical")
        self.assertEqual(v0_enriched["AI Usage Detected (Y/N)"], "Yes")
        self.assertIn("Larkspur", str(v0_enriched["Vendor Name"]))

    def test_9_automworx_not_classified_suspected_from_workload_automation_or_blog_title(self):
        """AutomWorx (V-001) must not be classified Suspected solely because workload automation or blog mentions AI Agents."""
        from src.assessment.classifier import NOT_DETECTED, classify_preliminary_ai_usage

        curated_ev = {
            "evidence_id": "E-CURATED-V1",
            "source_url": "https://www.automworx.com/services/workload-automation",
            "source_title": "[Curated Reference] AutomWorx Enterprise Workload Automation Services",
            "data_origin": "curated_case_study_disclosure",
            "excerpt": "Specialized consulting and managed support for Automic, AppWorx, and UC4 workload automation engines.",
            "ai_usage_detected": "No",
            "signal_types": ["non_ai_service_disclosure", "curated_case_study_disclosure"],
            "grade": {
                "score": 5,
                "source_type": "official_product_documentation",
                "directness": "indirect",
                "service_relevance": "relevant",
                "recency": "current",
            },
        }
        live_blog_ev = {
            "evidence_id": "E-LIVE-BLOG-V1",
            "source_url": "https://www.automworx.com/",
            "source_title": "AutomWorx Consulting",
            "data_origin": "live_public_web",
            "excerpt": "October 8, 2026 Before You Put AI Agents on the Schedule: Start With What Already Runs",
            "signal_types": ["generative_ai"],
            "grade": {
                "score": 2,
                "source_type": "official_product_documentation",
                "directness": "unclear",
                "service_relevance": "relevant",
                "recency": "current",
            },
        }

        # With both curated non-AI disclosure and blog commentary:
        ass = classify_preliminary_ai_usage([curated_ev, live_blog_ev], collection_complete=True)
        self.assertEqual(ass.status, NOT_DETECTED)
        self.assertNotIn("Suspected", ass.status)
        self.assertTrue(any("rule-based" in r.lower() or "curated" in r.lower() for r in ass.rationale))

    def test_10_unrelated_github_repositories_excluded_from_evidence_and_conclusions(self):
        """Unrelated third-party GitHub repositories (api-evangelist, bnycftc) are excluded from evidence register."""
        from src.collection.investigation import run_full_investigation
        from src.collection.fetcher import FetchResult

        class MockFetcher:
            def fetch_page(self, url):
                if "api-evangelist" in url:
                    html = "<html><body><h1>This is not our API. Independent profile by API Evangelist.</h1></body></html>"
                elif "bnycftc" in url:
                    html = "<html><body><h1>Shorts Creator app</h1></body></html>"
                else:
                    html = "<html><body><p>BNY Enterprise AI Hub uses Microsoft Azure OpenAI Service.</p></body></html>"
                return FetchResult(
                    url=url,
                    status="success",
                    retrieved_at="2026-10-13T00:00:00Z",
                    http_status=200,
                    html=html,
                )

        mock_searcher = lambda query: ([
            "https://github.com/api-evangelist/bny-mellon",
            "https://github.com/bnycftc/bnycftc.github.io",
        ], "GitHub public repository search returned 2 candidate(s).")

        res = run_full_investigation(
            vendor_id="V-005",
            vendor_name="BNY",
            seed_urls=["https://www.bny.com"],
            fetcher=MockFetcher(),
            github_searcher=mock_searcher,
            include_github=True,
            max_pages=5,
            max_depth=1,
        )

        evidence_urls = [e.get("source_url") for e in res["evidence_records"]]
        self.assertNotIn("https://github.com/api-evangelist/bny-mellon", evidence_urls)
        self.assertNotIn("https://github.com/bnycftc/bnycftc.github.io", evidence_urls)

        # Check collection log records exclusions
        excluded_log = [l for l in res["collection_log"] if l.get("status") == "excluded_unrelated_third_party"]
        self.assertEqual(len(excluded_log), 2)

    def test_11_evidence_description_provenance_labels_curated_vs_live(self):
        """Curated reference disclosures must be labeled as simulation references without misleading live retrieval dates."""
        enriched = self.batch_out["enriched_inventory_df"]
        v1_row = enriched[enriched["Vendor ID"] == "V-001"].iloc[0]
        v1_desc = str(v1_row.get("Evidence Description along with OSINT Source(s) / Evidence Link", ""))
        self.assertIn("[Curated Case-Study Reference — Not Independently Verified Live OSINT]", v1_desc)
        self.assertNotIn("Independently retrieved live OSINT — Source: [Curated Reference]", v1_desc)

    def test_12_unreviewed_evidence_status_is_needs_review_not_human_approved(self):
        """Automated evidence must remain 'Needs review' and preliminary risk findings 'Preliminary' until human action."""
        for ev in self.batch_out["all_evidence_records"]:
            self.assertEqual(ev.get("review_status"), "Needs review")
            self.assertNotEqual(ev.get("review_status"), "Accepted")

        for finding in self.batch_out["all_risk_findings"]:
            self.assertEqual(finding.get("status"), "Preliminary")
            self.assertNotEqual(finding.get("status"), "Reviewed")


    def test_13_automworx_rationale_does_not_claim_proof_of_no_ai_and_confidence_is_moderate(self):
        """AutomWorx rationale must not describe absence of evidence as proof of no AI deployment; confidence is Moderate."""
        inv = self.batch_out["investigation_assessments"].get("V-001")
        self.assertIsNotNone(inv)
        self.assertEqual(inv["confidence"], "Moderate")
        summary_text = str(inv["summary"]).lower()
        self.assertIn("absence of public evidence", summary_text)
        self.assertIn("does not", summary_text)
        # Check rationale
        combined_rat = " ".join(str(r) for r in inv.get("rationale", [])).lower()
        self.assertIn("absence of public evidence does not prove", combined_rat)

    def test_14_terrapin_blocked_live_collection_reflected_in_completeness_and_rationale(self):
        """Terrapin blocked live collection must be reflected in assessment completeness and rationale."""
        from src.collection.pipeline import assess_single_vendor
        row_v4 = self.df[self.df["Vendor ID"] == "V-004"].iloc[0]
        res = assess_single_vendor(row_v4, run_live_osint=True)
        self.assertEqual(res["ai_assessment"].status, "Inconclusive/Needs Review")
        summary = str(res["ai_assessment"].summary).lower()
        self.assertTrue("blocked by robots.txt" in summary or "robots" in summary)
        ev_desc = str(res["assessment_fields"]["Evidence Description along with OSINT Source(s) / Evidence Link"])
        self.assertIn("Completeness note: Live public web collection was blocked by vendor robots.txt policy restrictions", ev_desc)

    def test_15_fssi_and_terrapin_risk_rationales_maintain_conditional_claims(self):
        """FSSI and Terrapin risk rationales must maintain conditional claims without unbacked assumptions."""
        from src.collection.pipeline import CURATED_CASE_STUDY_DISCLOSURES
        for vid in ["V-003", "V-004"]:
            curated = CURATED_CASE_STUDY_DISCLOSURES[vid]
            rat = str(curated["risk_rationale"])
            self.assertIn("If ", rat)
            self.assertIn("Conversely, if ", rat)
            self.assertIn("unconfirmed pending", rat)
            self.assertNotIn("automated processing is largely deterministic", rat.lower())
            self.assertNotIn("suggests internal heuristics", rat.lower())

    def test_16_full_risk_findings_export_csv_synchronization(self):
        """Full risk findings register exports all findings across vendors and agrees with workbook Risk Findings sheet."""
        from src.assessment.risk import findings_as_records
        from src.ui.report_export import build_assessment_workbook
        from openpyxl import load_workbook
        from io import BytesIO

        findings = self.batch_out["all_risk_findings"]
        findings_df = pd.DataFrame(findings_as_records(findings))
        self.assertGreaterEqual(len(findings_df), 6)
        vendor_ids = set(findings_df["Vendor ID"].unique())
        self.assertTrue({"V-001", "V-002", "V-003", "V-004", "V-005", "V-006"}.issubset(vendor_ids))

        # Check workbook
        state = {
            "inventory_df": self.batch_out["enriched_inventory_df"],
            "evidence_records": self.batch_out["all_evidence_records"],
            "review_log": [],
            "risk_findings": findings,
            "collection_log": self.batch_out["all_collection_log"],
            "criticality_ratings": self.batch_out["criticality_ratings"],
            "investigation_assessments": self.batch_out["investigation_assessments"],
            "relationship_records": self.batch_out["all_relationship_records"],
        }
        with patch("streamlit.session_state", state):
            wb_bytes = build_assessment_workbook()
        wb = load_workbook(BytesIO(wb_bytes.getvalue()), read_only=True)
        rf_sheet = wb["Risk Findings"]
        self.assertEqual(rf_sheet.max_row - 1, len(findings_df))
        wb.close()

    def test_17_false_positive_ai_signals_rejected(self):
        """Evidence without qualifying AI signals or relevant service claims must not produce Confirmed or Suspected."""
        from src.assessment.classifier import _has_qualifying_ai_signals, classify_ai_usage

        # Non-AI IT workload/database evidence
        db_evidence = {
            "evidence_id": "E-DB-1",
            "vendor_id": "V-999",
            "source_url": "https://example.com/infrastructure",
            "source_title": "Enterprise Database Replication",
            "excerpt": "Enterprise PostgreSQL database replication and high-availability backup clustering.",
            "grade": {
                "score": 5,
                "directness": "direct",
                "service_relevance": "relevant",
                "recency": "current",
            },
            "review_status": "Accepted",
        }
        self.assertFalse(_has_qualifying_ai_signals(db_evidence))
        ass = classify_ai_usage([db_evidence], collection_complete=True)
        self.assertEqual(ass.status, "Not Detected")

        # Irrelevant evidence
        irrelevant_ai = {
            "evidence_id": "E-IRR-1",
            "vendor_id": "V-999",
            "source_url": "https://example.com/cafeteria",
            "source_title": "Cafeteria Smart Robot",
            "excerpt": "Our cafeteria features an AI robot for automated dish cleaning.",
            "grade": {
                "score": 4,
                "directness": "direct",
                "service_relevance": "irrelevant",
                "recency": "current",
            },
            "review_status": "Accepted",
        }
        self.assertFalse(_has_qualifying_ai_signals(irrelevant_ai))
        ass_irr = classify_ai_usage([irrelevant_ai], collection_complete=True)
        self.assertEqual(ass_irr.status, "Not Detected")

    def test_18_unrelated_accepted_evidence_does_not_validate_risk(self):
        """Accepted evidence only supports a risk finding when its content substantiates that finding's theme."""
        from src.assessment.risk import calculate_ai_cyber_risk_score, create_risk_finding

        finding = create_risk_finding(
            vendor_id="V-100",
            vendor_name="Acme Analytics",
            theme="retention_and_training",
            severity="Critical",
            likelihood="Likely",
            rationale="Customer data processed by external AI could be retained or used for foundation model training.",
            evidence_ids=["E-MENU"],
            status="Preliminary",
        ).to_dict()

        # Accepted evidence, but completely unrelated to retention_and_training
        unrelated_evidence = {
            "evidence_id": "E-MENU",
            "vendor_id": "V-100",
            "review_status": "Accepted",
            "source_title": "Office Cafeteria Guidelines",
            "excerpt": "Our building cafeteria offers organic meal subscriptions and daily lunch menus.",
            "grade": {
                "score": 4,
                "service_relevance": "relevant",
                "directness": "direct",
            },
        }

        eval_unrelated = calculate_ai_cyber_risk_score([finding], evidence_records=[unrelated_evidence])
        self.assertIsNone(eval_unrelated["score"])
        self.assertEqual(eval_unrelated["unsupported_findings"], 1)

        # Substantiating evidence
        substantiating_evidence = {
            "evidence_id": "E-MENU",
            "vendor_id": "V-100",
            "review_status": "Accepted",
            "source_title": "AI Model Training & Retention Policy",
            "excerpt": "Customer inputs and training data are stored for continuous machine learning model training.",
            "signal_types": ["model_training", "data_retention"],
            "grade": {
                "score": 5,
                "service_relevance": "relevant",
                "directness": "direct",
            },
        }
        eval_supported = calculate_ai_cyber_risk_score([finding], evidence_records=[substantiating_evidence])
        self.assertEqual(eval_supported["score"], 100.0)

    def test_19_incomplete_collection_preserves_uncertainty_in_sync(self):
        """sync_vendor_after_review must preserve Inconclusive when collection failed or was incomplete."""
        from src.assessment.review import sync_vendor_after_review

        # Vendor with no accepted positive evidence and a failed/blocked collection log
        collection_log = [{
            "vendor_id": "V-BLOCKED",
            "status": "blocked",
            "details": "Robots.txt policy disallowed scraping /platform",
        }]

        res = sync_vendor_after_review(
            "V-BLOCKED",
            evidence_records=[],
            collection_log=collection_log,
            collection_complete=False,
            collection_failed=True,
            failure_reason="Robots.txt policy disallowed scraping /platform",
        )
        self.assertEqual(res["ai_status"], "Inconclusive/Needs Review")
        self.assertEqual(res["detected_flag"], "Inconclusive")

    def test_20_reviewer_acceptance_distinct_from_automated_grading(self):
        """Automated grading leaves evidence in 'Needs review'; evidence is Accepted only through explicit action."""
        from src.assessment.evidence import create_evidence_record
        from src.assessment.review import review_evidence, apply_review_status

        candidate = {
            "source_url": "https://example.com/ai-docs",
            "source_title": "AI Documentation",
            "excerpt": "Our platform integrates machine learning models for predictive scoring.",
            "signal_types": ["machine_learning"],
            "matched_terms": ["machine learning"],
        }
        rec = create_evidence_record(
            candidate=candidate,
            vendor_id="V-777",
            vendor_name="Test Vendor",
            source_type="official_product_documentation",
            directness="direct",
            service_relevance="relevant",
        ).to_dict()

        # Automated pipeline produces 'Needs review'
        self.assertEqual(rec.get("review_status"), "Needs review")
        self.assertNotEqual(rec.get("review_status"), "Accepted")

        # Explicit reviewer decision transitions status to Accepted
        decision = review_evidence(
            rec,
            decision="Accepted",
            reviewer="Security Analyst",
            reason="Relevant and credible",
            notes="Verified documentation link.",
        )
        updated_rec = apply_review_status(rec, decision)
        self.assertEqual(updated_rec.get("review_status"), "Accepted")
        self.assertEqual(updated_rec.get("reviewer"), "Security Analyst")

    def test_21_unrelated_vendor_inventory_does_not_apply_case_study_fixtures(self):
        """Curated case-study fixtures are not applied to unrelated vendor inventories matching IDs by accident."""
        from src.collection.pipeline import assess_single_vendor
        from src.collection.fixtures import get_curated_case_study_disclosure

        # V-001 but named Acme Corp (unrelated to AutomWorx)
        unrelated_row = {
            "Vendor ID": "V-001",
            "Vendor Name": "Acme Logistics Corp",
            "Category": "Freight & Shipping",
            "Public Website": "www.acmelogistics.example",
            "Service / Product Provided": "Freight delivery and freight forwarding",
            "Operational Dependency": "Low",
            "Data Classification Accessed": "Internal operational shipping addresses.",
            "Data Volume (annual)": "50,000 shipments",
        }

        # Verify fixture getter returns None for unrelated vendor
        self.assertIsNone(get_curated_case_study_disclosure("V-001", vendor_name="Acme Logistics Corp"))

        # Verify assess_single_vendor does not apply AutomWorx curated fixture
        res = assess_single_vendor(unrelated_row, run_live_osint=False)
        self.assertNotIn("AutomWorx", str(res["assessment_fields"]["How Vendor May Be Using AI"]))
        self.assertNotIn("AppWorx", str(res["assessment_fields"]["Evidence Description along with OSINT Source(s) / Evidence Link"]))
        self.assertEqual(len(res["evidence_records"]), 0)


    def test_22_unsupported_baseline_findings_do_not_generate_score_and_distinguish_unassessed(self):
        """Unsupported baseline findings return score=None, status='No validated AI-specific risk findings', not zero."""
        from src.assessment.risk import calculate_ai_cyber_risk_score
        from src.ui.report_export import build_vendor_scores_dataframe

        # 1. AutomWorx baseline finding with rated severity/likelihood but no accepted positive evidence
        automworx_findings = [{
            "finding_id": "RISK-V-001-BASELINE-MONITORING",
            "vendor_id": "V-001",
            "theme": "evidence_gaps",
            "risk_type": "assurance_gap",
            "severity": "Low",
            "likelihood": "Unlikely",
            "evidence_ids": ["E-V001-01"],
            "title": "Baseline vendor AI deployment monitoring",
            "description": "Maintain periodic assurance monitoring for automated workloads.",
        }]
        # Curated disclosure indicating non-AI / rule-based workload automation
        curated_evidence = [{
            "evidence_id": "E-V001-01",
            "vendor_id": "V-001",
            "review_status": "Accepted",
            "signal_types": ["rule_based_automation"],
            "excerpt": "AppWorx workload automation platform provides rule-based scheduling without AI components.",
        }]

        risk_res = calculate_ai_cyber_risk_score(
            automworx_findings,
            evidence_records=curated_evidence,
        )

        # Must NOT manufacture a numeric score (e.g. 0.0)
        self.assertIsNone(risk_res["score"])
        self.assertEqual(risk_res["status"], "No validated AI-specific risk findings")
        self.assertEqual(risk_res["scored_findings"], 0)
        self.assertEqual(risk_res["unsupported_findings"], 1)
        self.assertEqual(risk_res["unassessed_findings"], 0)
        # Rationale must explicitly state unscored is not verified zero risk
        self.assertIn("verified zero-risk result", risk_res["rationale"].lower())
        self.assertIn("baseline assurance recommendations remain active", risk_res["rationale"].lower())

        # 2. Incomplete / unassessed findings must be clearly distinguished
        incomplete_findings = [{
            "finding_id": "RISK-V-INCOMPLETE-01",
            "vendor_id": "V-999",
            "theme": "sensitive_data_exposure",
            "risk_type": "conditional",
            "severity": None,
            "likelihood": "Possible",
            "evidence_ids": [],
        }]
        incomplete_res = calculate_ai_cyber_risk_score(incomplete_findings)
        self.assertIsNone(incomplete_res["score"])
        self.assertEqual(incomplete_res["status"], "Assessment incomplete / unassessed — reviewer ratings required")
        self.assertNotEqual(incomplete_res["status"], risk_res["status"])

        # 3. Export scores DataFrame separates metrics and shows distinct status
        import pandas as pd
        inventory_df = pd.DataFrame([{
            "Vendor ID": "V-001",
            "Vendor Name": "AutomWorx Consulting",
            "Service / Product Provided": "Workload automation",
        }])
        criticality_ratings = {
            "V-001": {
                "score": 62.5,
                "tier": "Tier 2 - High",
                "status": "Complete preliminary assessment",
                "factor_coverage": 1.0,
            }
        }
        investigation_assessments = {
            "V-001": {
                "status": "No",
                "confidence": "Medium",
                "rationale": ["Curated disclosure confirms deterministic rule-based scheduling."],
            }
        }
        scores_df = build_vendor_scores_dataframe(
            inventory_df,
            criticality_ratings=criticality_ratings,
            investigation_assessments=investigation_assessments,
            risk_findings=automworx_findings,
            evidence_records=curated_evidence,
        )
        v1_row = scores_df[scores_df["Vendor ID"] == "V-001"].iloc[0]
        # Separate metrics
        self.assertEqual(v1_row["Business Criticality (0-100)"], 62.5)
        self.assertEqual(v1_row["AI-use Classification"], "No")
        self.assertTrue(pd.isna(v1_row["AI Cybersecurity Risk (0-100)"]))
        self.assertEqual(v1_row["Cyber Risk Status"], "No validated AI-specific risk findings")


if __name__ == "__main__":
    unittest.main(verbosity=2)


